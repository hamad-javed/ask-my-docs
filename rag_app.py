import argparse

from docx import Document as DocxFile
from langchain_chroma import Chroma
from langchain_core.documents import Document
from langchain_core.output_parsers import StrOutputParser
from langchain_core.prompts import ChatPromptTemplate
from langchain_core.runnables import RunnablePassthrough
from langchain_ollama import OllamaEmbeddings, OllamaLLM
from langchain_text_splitters import RecursiveCharacterTextSplitter

import config


def read_file(path):
    """Return the text of a .txt or .docx file."""
    if path.lower().endswith(".docx"):
        doc = DocxFile(path)
        lines = [p.text for p in doc.paragraphs if p.text.strip()]
        # Resumes often put skills, dates or contact details in tables
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    lines.append(" | ".join(dict.fromkeys(cells)))
        return "\n".join(lines)
    with open(path, encoding="utf-8") as f:
        return f.read()


def format_docs(docs):
    return "\n\n".join(doc.page_content for doc in docs)


parser = argparse.ArgumentParser(description="Ask questions about a .txt or .docx file using local RAG.")
parser.add_argument("question", nargs="*", help="ask one question and exit (omit for interactive mode)")
parser.add_argument("-f", "--file", default=config.DEFAULT_SOURCE_FILE, help="document to load (.txt or .docx)")
args = parser.parse_args()

# 1. LOAD & CHUNK: Read the file and cut it into smaller pieces
docs = [Document(page_content=read_file(args.file), metadata={"source": args.file})]
text_splitter = RecursiveCharacterTextSplitter(
    chunk_size=config.CHUNK_SIZE, chunk_overlap=config.CHUNK_OVERLAP
)
chunks = text_splitter.split_documents(docs)

# 2. EMBED & STORE: Turn each chunk into an embedding and save it in Chroma.
# The DB is saved in ./chroma_db so it can be inspected (see view_db.py) and
# is rebuilt on every run so edits to the document are picked up.
embeddings = OllamaEmbeddings(model=config.EMBEDDING_MODEL)
vector_store = Chroma(
    collection_name=config.COLLECTION_NAME,
    embedding_function=embeddings,
    persist_directory=config.DB_DIRECTORY,
)
vector_store.reset_collection()
vector_store.add_documents(chunks)

# 3. CONFIGURE RETRIEVER: Fetch the chunks closest in meaning to the question
retriever = vector_store.as_retriever(search_kwargs={"k": config.TOP_K})

# 4. CHOOSE THE BRAIN: The LLM that writes the answer
llm = OllamaLLM(model=config.LLM_MODEL)

# 5. DESIGN THE PROMPT: Tell the LLM how to behave with the context
system_prompt = (
    "You are a helpful assistant answering questions about a document. "
    "Use only the retrieved context below. If something is not mentioned in the "
    "context, say clearly that the document does not mention it. Keep answers "
    "short and do not suggest other sources.\n\n"
    "{context}"
)
prompt = ChatPromptTemplate.from_messages([
    ("system", system_prompt),
    ("human", "{input}"),
])

# 6. PIPELINE CHAIN: question -> retriever -> prompt -> LLM -> text
rag_chain = (
    {"context": retriever | format_docs, "input": RunnablePassthrough()}
    | prompt
    | llm
    | StrOutputParser()
)

# 7. RUN THE QUERY: Answer a question passed on the command line,
# otherwise keep asking for questions until the user types "quit"
if args.question:
    user_question = " ".join(args.question)
    print(f"Question: {user_question}")
    print(f"Answer: {rag_chain.invoke(user_question)}")
else:
    print(f'Loaded {args.file} ({len(chunks)} chunks). Ask a question (type "quit" to exit).')
    while True:
        try:
            user_question = input("\nQuestion: ").strip()
        except (EOFError, KeyboardInterrupt):
            print()
            break
        if user_question.lower() in ("quit", "exit", "q"):
            break
        if not user_question:
            continue
        print(f"Answer: {rag_chain.invoke(user_question)}")
